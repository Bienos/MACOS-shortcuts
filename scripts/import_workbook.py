#!/usr/bin/env python3
"""Import the workbook's canonical shortcut table and curated 30-item path."""
from __future__ import annotations

import json
from pathlib import Path
from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1]
WORKBOOK_NAME = "Mac_Keyboard_Shortcuts_Master.xlsx"
WORKBOOK = APP / WORKBOOK_NAME
if not WORKBOOK.exists():
    WORKBOOK = APP.parent / WORKBOOK_NAME
OUT = APP / "data" / "shortcuts.json"

FIELDS = [
    "id", "category", "subcategory", "shortcut", "shortcutText", "function",
    "description", "context", "level", "usefulness", "frequency", "enabledByDefault",
    "configuration", "versionDependencies", "keyboardLayout", "notes", "verification",
    "sourceUrl", "checked", "sourceLearningStatus",
]

def clean(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value

def playground_for(record):
    """Map each source row by its actual action and source context."""
    category = str(record.get("category", "")).casefold()
    subcategory = str(record.get("subcategory", "")).casefold()
    context = str(record.get("context", "")).casefold()
    function = str(record.get("function", "")).casefold()
    if any(term in function for term in ("wylog", "wymuś zakończenie", "zakończ aplikację", "force quit", "uśpij mac", "wymuś ponowne uruchomienie", "wyłącz mac", "zamknij wszystkie aplikacje i wyłącz", "power")):
        return "guided", "Ryzykowna czynność: wizualizacja objaśnia skrót, ale nigdy nie wylogowuje, nie wyłącza ani nie zamyka aplikacji na urządzeniu."
    if any(term in function for term in ("zrzut", "zrzuty", "przechwyć ekran", "nagrywania ekranu", "screen capture", "screenshot")) or "zrzut" in category or "przechwytywanie" in subcategory:
        return "capture", "Scena pulpitu pozwala wybrać obszar lub okno i obejrzeć symulowany podgląd zrzutu."
    if "terminal" in category or "programowanie" in category or "terminal" in context or any(x in context for x in ("visual studio code", "xcode")):
        return "code", "Izolowany edytor lub terminal pokazuje zmianę w makiecie; żadne polecenie nie jest uruchamiane."
    if "przeglądark" in category or any(x in context for x in ("safari", "chrome", "firefox", "edge", "przeglądark")):
        return "browser", "Makieta przeglądarki pokazuje karty, pole adresu i lokalną historię symulacji."
    if "finder" in category or "pliki i foldery" in subcategory or "finder" in context or "finderze" in function or "zaznaczonego pliku" in function or "podgląd pliku" in function:
        return "files", "Makieta Findera pozwala zaznaczać, otwierać i zmieniać położenie przykładowych plików."
    if any(term in function for term in ("ukryj aplikację", "ukryj pozostałe", "minimalizuj", "pełny ekran", "przełącz do następnej aplikacji", "przełącz okna", "nowe okno", "zamknij okno", "zamknij wszystkie okna", "mission control", "spaces", "pokaż biurko", "nowa karta")):
        return "windows", "Makieta okien i kart odzwierciedla czynność z rekordu; zmiany pozostają w tej stronie."
    if "system" in category or "dostępność" in category or "zasilanie" in subcategory or "dostępność" in subcategory or any(x in function for x in ("spotlight", "podgląd znaków", "odwróć kolory", "voiceover", "zoom", "dyktowanie", "zasilania", "funkcje dostępności")):
        return "guided", "Bezpieczna wizualizacja wyjaśnia działanie bez wylogowania, zmiany ustawień ani sterowania macOS."
    if "tekst" in category or "dokumenty i edycja" in subcategory or "nawigacja kursora" in subcategory or "zaznaczanie" in subcategory or "formatowanie" in subcategory or (category == "podstawowe" and any(x in function for x in ("wytnij", "kopiuj", "wklej", "cofnij", "ponów", "zaznacz", "znajdź", "otwórz", "drukuj", "zapisz"))):
        return "text", "Edytowalny dokument demonstracyjny pokazuje zaznaczenie lub zmianę tekstu."
    if any(x in subcategory for x in ("okna", "mission control", "interfejs", "sesje", "aplikacje")):
        return "windows", "Interaktywna makieta okien i kart pokazuje zmianę widoku bez sterowania prawdziwym systemem."
    if "aplikacje" in category:
        return "app", "Makieta aplikacji pokazuje stan i przykładową reakcję; rzeczywiste działanie zależy od kontekstu źródłowego."
    return "guided", "Guided visualization: workbook metadata is preserved, but the action is not described well enough for a confident behavior simulation."

def main():
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = wb["01_WSZYSTKIE"]
    rows = list(sheet.iter_rows(min_row=5, max_col=20, values_only=True))
    records = []
    for row in rows:
        if not row[0]:
            continue
        record = {key: clean(value) for key, value in zip(FIELDS, row)}
        record["playgroundType"], record["playgroundReason"] = playground_for(record)
        records.append(record)

    path_sheet = wb["22_UCZ_SIE_30"]
    path = []
    for row in path_sheet.iter_rows(min_row=5, max_col=9, values_only=True):
        if not row[2]:
            continue
        path.append({
            "order": clean(row[0]), "stage": clean(row[1]), "id": str(row[2]).zfill(4),
            "shortcut": clean(row[3]), "function": clean(row[4]), "context": clean(row[5]),
            "example": clean(row[6]), "sourceStatus": clean(row[7]), "sourceUrl": clean(row[8]),
        })

    legends = []
    legend_sheet = wb["17_KLAWIATURA"]
    for row in legend_sheet.iter_rows(min_row=5, max_col=3, values_only=True):
        if row[0]:
            legends.append({"symbol": clean(row[0]), "name": clean(row[1]), "meaning": clean(row[2])})

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({
        "sourceWorkbook": WORKBOOK.name,
        "sourceSheet": sheet.title,
        "recordCount": len(records),
        "pathSourceSheet": path_sheet.title,
        "records": records,
        "learningPath": path,
        "keyLegend": legends,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Imported {len(records)} shortcuts and {len(path)} curated path entries to {OUT}")

if __name__ == "__main__":
    main()
